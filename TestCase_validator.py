#!/usr/bin/env python3
"""
ShieldDesk Combined Test Validator (Static Code Check + Live Browser Run)
============================================================================

This runs BOTH earlier tools together and cross-checks their answers
against each other, instead of trusting either one alone:

  1. STATIC CHECK  - searches the source code inside your ZIP for literal
                      text from each test case (error messages, button
                      labels, emails). No browser, no running app needed
                      for this half.
  2. LIVE RUN       - actually opens a real Chromium browser, logs in,
                      clicks through the test's steps against your running
                      app, and checks what's really on the page.

WHY COMBINE THEM
------------------
Each one is wrong in a different, predictable way:
  - The static check can say "no evidence" for a feature that's real but
    worded differently in code (a false FAIL).
  - The live run can say "couldn't confirm" for a feature that's real but
    whose button/field it couldn't find with a guessed selector (a false
    REVIEW/FAIL).

When both checks agree, that's strong evidence either way. When they
disagree, that disagreement itself is useful information -- it usually
means either a genuine bug, or an automation limitation on one side, and
the evidence log tells you which. The combined Status you get is
therefore more trustworthy than either check run alone, and every verdict
comes with a Confidence rating (High/Medium/Low) so you know how much to
trust it.

    Static \\ Live   PASS              REVIEW                 FAIL
    -----------------------------------------------------------------------
    PASS            PASS (High)       REVIEW (Medium)        REVIEW (Medium)
    REVIEW          PASS (Medium)     REVIEW (Low)            FAIL (Medium)
    FAIL            PASS (Medium)     REVIEW (Medium)        FAIL (High)

(Static PASS = literal text found in code. Static FAIL = nothing found.
Live PASS = browser confirmed the expected result on screen. See each
engine's own notes below for what REVIEW means on that side.)

OUTPUT
------
Same 12 columns as your original workbook -- Status is overwritten with
the combined verdict, nothing added or removed from that sheet. A new
'Automated Evidence Log' sheet holds both engines' reasoning side by side
(static evidence + file:line, live steps executed/skipped, screenshot
path) plus the Confidence rating and why the two engines agreed or
disagreed.

SETUP
-----
    pip install openpyxl playwright --break-system-packages
    playwright install chromium

USAGE
-----
Both checks (default, recommended):

    python shielddesk_combined_validator.py \
        --excel ShieldDesk_Final_Test_Case_Suite.xlsx \
        --zip   shielddesk_source.zip \
        --base-url http://localhost:3000

Only the static code check (no running app / Playwright needed):

    python shielddesk_combined_validator.py \
        --excel ShieldDesk_Final_Test_Case_Suite.xlsx \
        --zip   shielddesk_source.zip \
        --static-only

Only the live browser run (same as the standalone live runner):

    python shielddesk_combined_validator.py \
        --excel ShieldDesk_Final_Test_Case_Suite.xlsx \
        --zip   shielddesk_source.zip \
        --base-url http://localhost:3000 \
        --live-only

Let the script start your app itself, watch it run, and see per-step
selector attempts while you dial in accuracy:

    python shielddesk_combined_validator.py \
        --excel ShieldDesk_Final_Test_Case_Suite.xlsx \
        --zip   shielddesk_source.zip \
        --base-url http://localhost:3000 \
        --start-cmd "npm install && npm start" \
        --max-tests 10 --headed --debug

--selector-map (JSON of CSS selector overrides for fields/buttons) works
the same as in the standalone live runner -- see that script's docstring
or run with --debug to see what it's trying to match on your real app.
"""

import argparse
import json
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
import zipfile
from collections import Counter
from urllib.parse import urljoin

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("This script needs openpyxl. Install it with:\n"
              "    pip install openpyxl --break-system-packages")


# ============================================================================
# Shared configuration
# ============================================================================

EVIDENCE_SHEET_NAME = "Automated Evidence Log"
EVIDENCE_HEADERS = [
    "Test Case ID", "Previous Status", "Final Status", "Confidence",
    "Static Result", "Static Evidence (file:line)",
    "Live Result", "Live Steps (parsed/executed/skipped)", "Live Notes", "Screenshot",
]

QUOTE_RE = re.compile(r"'([^'\n]{2,120})'|\"([^\"\n]{2,120})\"")
REQUIRED_COLUMNS = ["Test Case ID", "Module", "Test Case Title", "Preconditions",
                    "Test Steps", "Test Data", "Expected Result", "Role / Persona", "Status"]

STATUS_FILLS = {
    "PASS": PatternFill("solid", fgColor="C6EFCE"),
    "REVIEW": PatternFill("solid", fgColor="FFEB9C"),
    "FAIL": PatternFill("solid", fgColor="FFC7CE"),
}


def extract_zip(zip_path):
    tmpdir = tempfile.mkdtemp(prefix="tcv_combined_")
    try:
        with zipfile.ZipFile(zip_path) as zf:
            zf.extractall(tmpdir)
    except zipfile.BadZipFile:
        sys.exit("The file passed to --zip is not a valid ZIP archive.")
    return tmpdir


# ============================================================================
# ENGINE 1: Static source-code check
# ============================================================================

TEXT_EXTS = {
    ".py", ".js", ".jsx", ".ts", ".tsx", ".mjs", ".cjs",
    ".java", ".php", ".rb", ".go", ".cs", ".c", ".cpp", ".h", ".hpp",
    ".html", ".htm", ".css", ".scss", ".sass", ".vue", ".svelte",
    ".json", ".yml", ".yaml", ".xml", ".sql", ".md", ".txt", ".env",
    ".graphql", ".gql",
}
SKIP_DIRS = {
    "node_modules", ".git", ".hg", ".svn", "venv", ".venv", "env",
    "dist", "build", ".next", "__pycache__", "coverage", ".idea",
    ".vscode", "target", "vendor", ".pytest_cache", ".cache",
}
MAX_FILE_SIZE = 1_500_000
MAX_HITS_PER_TERM = 3
EMAIL_RE = re.compile(r"[\w.\-]+@[\w.\-]+\.\w+")


class CodebaseIndex:
    """Loads every text-like file under `root` and supports substring search
    with a per-term cache (the same Module/Role keywords repeat across
    hundreds of test cases)."""

    def __init__(self, root):
        self.files = {}
        self.blobs = {}
        self._term_cache = {}
        for dirpath, dirnames, filenames in os.walk(root):
            dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS and not d.startswith(".")]
            for fn in filenames:
                ext = os.path.splitext(fn)[1].lower()
                if ext not in TEXT_EXTS:
                    continue
                full = os.path.join(dirpath, fn)
                try:
                    if os.path.getsize(full) > MAX_FILE_SIZE:
                        continue
                    with open(full, "r", encoding="utf-8", errors="ignore") as fh:
                        text = fh.read()
                except OSError:
                    continue
                rel = os.path.relpath(full, root)
                self.files[rel] = text.splitlines()
                self.blobs[rel] = text.lower()

    def search(self, term, max_hits=MAX_HITS_PER_TERM):
        term = (term or "").strip()
        if len(term) < 3:
            return []
        key = term.lower()
        if key in self._term_cache:
            return self._term_cache[key]
        hits = []
        for rel, blob in self.blobs.items():
            if key not in blob:
                continue
            for i, line in enumerate(self.files[rel], start=1):
                if key in line.lower():
                    snippet = line.strip()
                    if len(snippet) > 140:
                        snippet = snippet[:140] + "..."
                    hits.append((rel, i, snippet))
                    if len(hits) >= max_hits:
                        break
            if len(hits) >= max_hits:
                break
        self._term_cache[key] = hits
        return hits


def extract_static_targets(row):
    text_fields = [
        row.get("Test Case Title") or "", row.get("Test Steps") or "",
        row.get("Expected Result") or "", row.get("Test Data") or "",
        row.get("Preconditions") or "",
    ]
    literals = set()
    for text in text_fields:
        for m in QUOTE_RE.finditer(text):
            val = (m.group(1) or m.group(2) or "").strip()
            if len(val) >= 3:
                literals.add(val)
        for m in EMAIL_RE.finditer(text):
            literals.add(m.group(0))

    keywords = set()
    module_role = f"{row.get('Module') or ''} {row.get('Role / Persona') or ''}"
    for chunk in re.split(r"[/,&()]", module_role):
        chunk = chunk.strip()
        if len(chunk) >= 4 and chunk.lower() not in {"any", "user", "users"}:
            keywords.add(chunk)
    return literals, keywords


def evaluate_static(index, row):
    """Returns (status, evidence_summary) where status is PASS/REVIEW/FAIL."""
    literals, keywords = extract_static_targets(row)
    literal_hits = {lit: index.search(lit) for lit in literals}
    literal_hits = {k: v for k, v in literal_hits.items() if v}
    keyword_hits = {kw: index.search(kw) for kw in keywords}
    keyword_hits = {k: v for k, v in keyword_hits.items() if v}

    if literal_hits:
        status = "PASS"
    elif keyword_hits:
        status = "REVIEW"
    else:
        status = "FAIL"

    evidence_bits = []
    source = literal_hits if literal_hits else keyword_hits
    for term, hits in list(source.items())[:3]:
        rel, line_no, _snippet = hits[0]
        evidence_bits.append(f"'{term}' -> {rel}:{line_no}")
    evidence = "; ".join(evidence_bits) if evidence_bits else "No matching code found"
    return status, evidence


# ============================================================================
# ENGINE 2: Live browser run (Playwright)
# ============================================================================

SELECTOR_PROBE_TIMEOUT = 2500
CREDS_RE = re.compile(r"email:\s*([^\s/]+)\s*/\s*password:\s*(\S+)", re.I)
LOGIN_AS_RE = re.compile(r"log\s*in\s+as\s+(?:an?\s+)?([A-Za-z][\w\s/]*?)(?:\s*[\(,\.]|\s+and\b|$)", re.I)
ENTER_RE = re.compile(r"(?:enter|type|fill\s+(?:in|out)?)\s+(?:an?\s+|the\s+)?([\w\s]{2,25}?)\s*(?:as\s+|with\s+)?['\"]([^'\"]+)['\"]", re.I)
CLICK_RE = re.compile(r"(?:click|tap|select|choose|press)\s+(?:on\s+)?(?:the\s+)?['\"]([^'\"]+)['\"]", re.I)
NAV_RE = re.compile(r"(?:navigate|go)\s+to\s+['\"]?([^'\"\n]+?)['\"]?\s*$", re.I)


def wait_for_server(url, timeout):
    deadline = time.time() + timeout
    last_err = None
    while time.time() < deadline:
        try:
            urllib.request.urlopen(url, timeout=3)
            return True
        except Exception as e:
            last_err = e
        time.sleep(1)
    print(f"      server did not respond at {url} within {timeout}s (last error: {last_err})")
    return False


def start_app(start_cmd, cwd, log_path):
    print(f"      running: {start_cmd}  (cwd={cwd})")
    logf = open(log_path, "w")
    kwargs = {}
    if os.name == "posix":
        kwargs["preexec_fn"] = os.setsid
    proc = subprocess.Popen(start_cmd, shell=True, cwd=cwd, stdout=logf, stderr=subprocess.STDOUT, **kwargs)
    return proc, logf


def stop_app(proc, logf):
    if proc is None:
        return
    try:
        if os.name == "posix":
            os.killpg(os.getpgid(proc.pid), signal.SIGTERM)
        else:
            proc.terminate()
        time.sleep(1)
        if proc.poll() is None:
            if os.name == "posix":
                os.killpg(os.getpgid(proc.pid), signal.SIGKILL)
            else:
                proc.kill()
    except Exception:
        pass
    finally:
        try:
            logf.close()
        except Exception:
            pass


def build_credentials(rows):
    creds = {}
    for row in rows:
        blob = f"{row.get('Test Data') or ''}\n{row.get('Test Steps') or ''}"
        m = CREDS_RE.search(blob)
        if not m:
            continue
        role = (row.get("Role / Persona") or "").strip()
        if not role:
            continue
        for part in re.split(r"[/,]", role):
            key = part.strip().lower()
            if key and key not in creds:
                creds[key] = (m.group(1), m.group(2))
    return creds


def find_creds(creds, role_text):
    role_text = (role_text or "").strip().lower()
    if not role_text:
        return None
    if role_text in creds:
        return creds[role_text]
    for key, val in creds.items():
        if key in role_text or role_text in key:
            return val
    return None


def parse_steps(steps_text):
    lines = [re.sub(r"^\s*\d+[\.\)]\s*", "", ln).strip()
             for ln in (steps_text or "").splitlines() if ln.strip()]
    actions = []
    for line in lines:
        m = LOGIN_AS_RE.search(line)
        if m:
            actions.append(("login_as", m.group(1).strip(), line))
            continue
        m = ENTER_RE.search(line)
        if m:
            actions.append(("enter", m.group(1).strip(), m.group(2).strip(), line))
            continue
        m = CLICK_RE.search(line)
        if m:
            actions.append(("click", m.group(1).strip(), line))
            continue
        m = NAV_RE.search(line)
        if m:
            actions.append(("goto", m.group(1).strip(), line))
            continue
        actions.append(("skip", line))
    return actions


def _fill_selector(page, sel, value, timeout, probe, debug, label):
    try:
        loc = page.locator(sel).first
        loc.wait_for(state="visible", timeout=probe)
        loc.scroll_into_view_if_needed(timeout=probe)
        loc.fill(value, timeout=timeout)
        if debug:
            print(f"        [fill] matched '{label}'")
        return True
    except Exception:
        return False


def fill_field(page, field_label, value, timeout, selector_map=None, debug=False):
    lf = field_label.lower()
    probe = min(timeout, SELECTOR_PROBE_TIMEOUT)
    if selector_map:
        for key, sel in selector_map.get("fields", {}).items():
            if key.lower() in lf or lf in key.lower():
                if _fill_selector(page, sel, value, timeout, probe, debug, f"map:{key}"):
                    return True

    selectors = []
    if "email" in lf or "e-mail" in lf or "username" in lf or "user name" in lf:
        selectors += ["input[type=email]", "input[name*=email i]", "input[id*=email i]",
                      "input[placeholder*=email i]", "input[aria-label*=email i]",
                      "input[name*=user i]", "input[id*=user i]"]
    elif "password" in lf or "pass" in lf:
        selectors += ["input[type=password]", "input[name*=pass i]", "input[id*=pass i]",
                      "input[placeholder*=pass i]", "input[aria-label*=pass i]"]
    else:
        safe = field_label.replace('"', '')
        selectors += [f'input[placeholder*="{safe}" i]', f'input[name*="{safe}" i]',
                      f'input[id*="{safe}" i]', f'input[aria-label*="{safe}" i]',
                      f'textarea[placeholder*="{safe}" i]', f'[data-testid*="{safe.lower()}"]']

    for sel in selectors:
        if _fill_selector(page, sel, value, timeout, probe, debug, sel):
            return True
    try:
        loc = page.get_by_label(field_label, exact=False).first
        loc.wait_for(state="visible", timeout=probe)
        loc.fill(value, timeout=timeout)
        if debug:
            print(f"        [fill] '{field_label}' -> matched via label")
        return True
    except Exception:
        if debug:
            print(f"        [fill] '{field_label}' -> NO SELECTOR MATCHED")
        return False


def _click_selector(page, sel, timeout, probe, debug, label):
    try:
        loc = page.locator(sel).first
        loc.wait_for(state="visible", timeout=probe)
        loc.click(timeout=timeout)
        if debug:
            print(f"        [click] matched '{label}'")
        return True
    except Exception:
        return False


def click_target(page, text, timeout, selector_map=None, debug=False):
    probe = min(timeout, SELECTOR_PROBE_TIMEOUT)
    if selector_map:
        for key, sel in selector_map.get("buttons", {}).items():
            if key.lower() == text.lower() or key.lower() in text.lower():
                if _click_selector(page, sel, timeout, probe, debug, f"map:{key}"):
                    return True

    escaped = re.escape(text)
    attempts = [
        (lambda: page.get_by_role("button", name=re.compile(escaped, re.I)).first, "role=button"),
        (lambda: page.get_by_role("link", name=re.compile(escaped, re.I)).first, "role=link"),
        (lambda: page.get_by_role("menuitem", name=re.compile(escaped, re.I)).first, "role=menuitem"),
        (lambda: page.get_by_role("tab", name=re.compile(escaped, re.I)).first, "role=tab"),
        (lambda: page.locator(f'[data-testid*="{text.lower()}" i]').first, "data-testid"),
        (lambda: page.get_by_text(text, exact=False).first, "text"),
    ]
    for get_loc, label in attempts:
        try:
            loc = get_loc()
            loc.wait_for(state="visible", timeout=probe)
            loc.scroll_into_view_if_needed(timeout=probe)
            loc.click(timeout=timeout)
            if debug:
                print(f"        [click] '{text}' -> matched via {label}")
            return True
        except Exception:
            continue
    if debug:
        print(f"        [click] '{text}' -> NO SELECTOR MATCHED")
    return False


def do_login(page, base_url, creds, timeout, settle_ms=0, selector_map=None, debug=False):
    if not creds:
        return False
    email, password = creds
    try:
        page.goto(base_url, timeout=timeout, wait_until="domcontentloaded")
        if settle_ms:
            page.wait_for_timeout(settle_ms)
    except Exception:
        return False
    ok_email = fill_field(page, "email", email, timeout, selector_map, debug)
    ok_pass = fill_field(page, "password", password, timeout, selector_map, debug)
    ok_click = (click_target(page, "Sign In", timeout, selector_map, debug)
                or click_target(page, "Log In", timeout, selector_map, debug)
                or click_target(page, "Login", timeout, selector_map, debug)
                or click_target(page, "Submit", timeout, selector_map, debug))
    if not ok_click:
        try:
            page.keyboard.press("Enter")
            ok_click = True
        except Exception:
            pass
    try:
        page.wait_for_load_state("networkidle", timeout=timeout)
    except Exception:
        pass
    if settle_ms:
        page.wait_for_timeout(settle_ms)
    return ok_email and ok_pass and ok_click


def precondition_role(row):
    m = LOGIN_AS_RE.search(row.get("Preconditions") or "") or \
        re.search(r"logged\s*in\s*as\s*(?:an?\s*)?([A-Za-z][\w\s/]*?)(?:[,\.\(]|$)",
                   row.get("Preconditions") or "", re.I)
    if m:
        return m.group(1).strip()
    return None


def _text_matches(literal, body_text):
    lit_lower = literal.lower()
    if lit_lower in body_text:
        return True
    norm = lambda s: re.sub(r"[^\w\s]", "", s).strip()
    norm_lit = norm(lit_lower)
    norm_body = norm(body_text)
    return bool(norm_lit) and norm_lit in norm_body


def run_live(browser, base_url, row, creds, timeout_ms, screenshots_dir,
             settle_ms=300, selector_map=None, debug=False):
    """Returns a dict describing what actually happened when a real browser
    ran this test case's steps against the live app."""
    test_id = row.get("Test Case ID") or "UNKNOWN"
    module = (row.get("Module") or "").strip().lower()
    steps = parse_steps(row.get("Test Steps"))
    if debug:
        print(f"      -- {test_id}: parsed {len(steps)} step(s): {steps}")

    context = browser.new_context()
    page = context.new_page()
    notes = []
    executed = skipped = 0
    bootstrap_failed = False
    literals = []
    found_text = None
    expected_found = None
    soft_note = None
    screenshot_path = ""
    total_steps = len(steps)

    try:
        has_login_action = any(a[0] == "login_as" for a in steps)
        pre_role = precondition_role(row)
        needs_bootstrap = (module != "login" and not has_login_action and
                            re.search(r"logged\s*in|authenticat", row.get("Preconditions") or "", re.I))
        if needs_bootstrap:
            role_text = pre_role or row.get("Role / Persona") or ""
            cred = find_creds(creds, role_text)
            if cred and do_login(page, base_url, cred, timeout_ms, settle_ms, selector_map, debug):
                notes.append(f"bootstrap login as '{role_text}' ok")
            else:
                bootstrap_failed = True
                notes.append(f"bootstrap login as '{role_text}' FAILED or no known credentials")
        elif not has_login_action:
            try:
                page.goto(base_url, timeout=timeout_ms, wait_until="domcontentloaded")
                if settle_ms:
                    page.wait_for_timeout(settle_ms)
            except Exception as e:
                notes.append(f"initial page load failed: {e}")

        for action in steps:
            kind = action[0]
            raw_line = action[-1]
            try:
                if kind == "login_as":
                    role_text = action[1]
                    cred = find_creds(creds, role_text)
                    if cred and do_login(page, base_url, cred, timeout_ms, settle_ms, selector_map, debug):
                        executed += 1
                    else:
                        skipped += 1
                        bootstrap_failed = True
                        notes.append(f"could not log in as '{role_text}'")
                elif kind == "enter":
                    field, value = action[1], action[2]
                    if fill_field(page, field, value, timeout_ms, selector_map, debug):
                        executed += 1
                    else:
                        skipped += 1
                        notes.append(f"could not find field for: {raw_line}")
                elif kind == "click":
                    target = action[1]
                    if click_target(page, target, timeout_ms, selector_map, debug):
                        executed += 1
                        if settle_ms:
                            page.wait_for_timeout(settle_ms)
                    else:
                        skipped += 1
                        notes.append(f"could not find element for: {raw_line}")
                elif kind == "goto":
                    target = action[1]
                    if click_target(page, target, timeout_ms, selector_map, debug):
                        executed += 1
                    else:
                        try:
                            page.goto(urljoin(base_url, target), timeout=timeout_ms)
                            executed += 1
                        except Exception:
                            skipped += 1
                            notes.append(f"could not navigate: {raw_line}")
                else:
                    skipped += 1
            except Exception as e:
                skipped += 1
                notes.append(f"error on step '{raw_line}': {e}")

        try:
            page.wait_for_load_state("networkidle", timeout=min(timeout_ms, 5000))
        except Exception:
            pass

        literals = [(m.group(1) or m.group(2)) for m in QUOTE_RE.finditer(row.get("Expected Result") or "")]
        if literals:
            try:
                body_text = page.locator("body").inner_text(timeout=timeout_ms).lower()
            except Exception:
                body_text = ""
            expected_found = all(_text_matches(lit, body_text) for lit in literals)
            found_text = "; ".join(literals)
        else:
            m = re.search(r"redirect(?:ed)?\s+to\s+(?:the\s+)?([A-Za-z][\w\s\-]{3,60}?dashboard)",
                           row.get("Expected Result") or "", re.I)
            if m:
                phrase = m.group(1)
                tokens = [t for t in re.findall(r"[A-Za-z]+", phrase)
                          if len(t) > 2 and t.lower() not in {"the", "and", "for", "user", "their"}]
                try:
                    body_text2 = page.locator("body").inner_text(timeout=timeout_ms).lower()
                except Exception:
                    body_text2 = ""
                url_changed = page.url.rstrip("/") != base_url.rstrip("/")
                hit = sum(1 for t in tokens if t.lower() in body_text2)
                soft_ok = bool(tokens) and url_changed and hit >= max(1, len(tokens) // 2)
                found_text = f"(soft) {phrase}"
                expected_found = soft_ok
                soft_note = f"soft redirect check on '{phrase}': {'matched' if soft_ok else 'no match'}"

        try:
            os.makedirs(screenshots_dir, exist_ok=True)
            screenshot_path = os.path.join(screenshots_dir, f"{test_id}.png")
            page.screenshot(path=screenshot_path, full_page=True)
        except Exception as e:
            notes.append(f"screenshot failed: {e}")

        had_action_lines = total_steps - sum(1 for a in steps if a[0] == "skip")
        automation_incomplete = (had_action_lines > 0 and skipped > 0)

        if bootstrap_failed:
            status = "REVIEW"
            notes.append("could not establish required login state -- result not trustworthy")
        elif automation_incomplete:
            status = "REVIEW"
        elif expected_found is not None:
            status = "PASS" if expected_found else "FAIL"
            if soft_note:
                notes.append(soft_note)
        else:
            status = "REVIEW"
            notes.append("Expected Result had no quoted literal or recognizable redirect phrase to verify")

    except Exception as e:
        status = "FAIL"
        notes.append(f"unhandled execution error: {e}")
    finally:
        context.close()

    return {
        "status": status,
        "steps_summary": f"{total_steps}/{executed}/{skipped}",
        "notes": "; ".join(notes) if notes else "",
        "screenshot": screenshot_path,
    }


# ============================================================================
# Combining the two verdicts
# ============================================================================

# (static_status, live_status) -> (final_status, confidence, note)
COMBINE_MATRIX = {
    ("PASS", "PASS"):   ("PASS",   "High",   "Both checks agree: code exists and browser confirmed it"),
    ("PASS", "REVIEW"): ("REVIEW", "Medium", "Code exists but the live run couldn't fully confirm -- check screenshot"),
    ("PASS", "FAIL"):   ("REVIEW", "Medium", "Code exists but live run did not see the expected result -- possible real bug or automation/selector issue"),
    ("REVIEW", "PASS"): ("PASS",   "Medium", "Live run confirmed the behavior even though static evidence was only partial"),
    ("REVIEW", "REVIEW"): ("REVIEW", "Low",  "Neither check found strong evidence either way -- needs a human look"),
    ("REVIEW", "FAIL"): ("FAIL",   "Medium", "Weak code evidence and the live run did not confirm it -- likely a real gap"),
    ("FAIL", "PASS"):   ("PASS",   "Medium", "Live run confirmed the behavior despite no matching text found in code (may be dynamic/generated content)"),
    ("FAIL", "REVIEW"): ("REVIEW", "Medium", "No code evidence and the live run couldn't confirm either -- treat as likely gap, verify manually"),
    ("FAIL", "FAIL"):   ("FAIL",   "High",   "Both checks agree: no supporting code and the live run did not see the expected result"),
}


def combine_verdicts(static_status, live_status):
    if static_status is None:
        return live_status, ("High" if live_status in ("PASS", "FAIL") else "Low"), "Live run only (static check skipped)"
    if live_status is None:
        return static_status, ("Medium" if static_status == "PASS" else "Low"), "Static check only (live run skipped)"
    return COMBINE_MATRIX[(static_status, live_status)]


# ============================================================================
# Workbook I/O
# ============================================================================

def load_header_map(ws):
    headers = [c.value for c in ws[1]]
    col_of = {h: i + 1 for i, h in enumerate(headers) if h}
    for name in REQUIRED_COLUMNS:
        if name not in col_of:
            sys.exit(f"Expected column '{name}' not found in sheet '{ws.title}'. Found: {headers}")
    return col_of


def write_evidence_log(wb, evidence_rows):
    if EVIDENCE_SHEET_NAME in wb.sheetnames:
        del wb[EVIDENCE_SHEET_NAME]
    ws = wb.create_sheet(EVIDENCE_SHEET_NAME)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    for i, name in enumerate(EVIDENCE_HEADERS, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = header_font
        cell.fill = header_fill
    for r, row in enumerate(evidence_rows, start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    widths = [18, 16, 14, 12, 14, 45, 14, 26, 45, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def write_summary_block(wb, counts, changed, total, zip_path, base_url, mode, n_files):
    if "Summary" not in wb.sheetnames:
        return
    ws = wb["Summary"]
    start = ws.max_row + 3
    ws.cell(row=start, column=2, value=f"COMBINED VALIDATION ({mode})").font = Font(bold=True, size=12)
    rows = [
        ("Source ZIP scanned", os.path.basename(zip_path)),
        ("Text files indexed", n_files if n_files is not None else "n/a (static check skipped)"),
        ("App URL tested", base_url or "n/a (live run skipped)"),
        ("Total test cases evaluated", total),
        ("Status = PASS", counts.get("PASS", 0)),
        ("Status = REVIEW", counts.get("REVIEW", 0)),
        ("Status = FAIL", counts.get("FAIL", 0)),
        ("Status changed from original", changed),
        ("Full evidence + confidence ratings", f"see '{EVIDENCE_SHEET_NAME}' sheet"),
    ]
    r = start + 1
    for label, val in rows:
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=3, value=val)
        r += 1
    note = ("Note: each row's Confidence rating (in the evidence log) tells you how much the two "
            "checks agreed. High confidence rows need little review; Medium/Low confidence rows "
            "are exactly where the static and live checks disagreed or were both weak -- look at "
            "those first.")
    ws.cell(row=r + 1, column=2, value=note).font = Font(italic=True, size=9, color="808080")


# ============================================================================
# Main
# ============================================================================

def main():
    p = argparse.ArgumentParser(description="Validate test cases with both a static code check and a live browser run.")
    p.add_argument("--excel", required=True)
    p.add_argument("--zip", required=True)
    p.add_argument("--base-url", default=None, help="Required unless --static-only")
    p.add_argument("--start-cmd", default=None)
    p.add_argument("--cwd", default=None)
    p.add_argument("--output", default=None)
    p.add_argument("--sheet", default="Test Cases")
    p.add_argument("--static-only", action="store_true", help="Skip the live browser run entirely")
    p.add_argument("--live-only", action="store_true", help="Skip the static code check entirely")
    p.add_argument("--startup-timeout", type=int, default=60)
    p.add_argument("--per-test-timeout", type=int, default=20000, help="ms")
    p.add_argument("--settle-ms", type=int, default=400)
    p.add_argument("--selector-map", default=None)
    p.add_argument("--debug", action="store_true")
    p.add_argument("--screenshots-dir", default="screenshots")
    p.add_argument("--headed", action="store_true")
    p.add_argument("--keep-server", action="store_true")
    p.add_argument("--max-tests", type=int, default=None)
    p.add_argument("--filter-module", default=None)
    args = p.parse_args()

    if args.static_only and args.live_only:
        sys.exit("--static-only and --live-only can't both be set.")

    do_static = not args.live_only
    do_live = not args.static_only

    if not os.path.isfile(args.excel):
        sys.exit(f"Excel file not found: {args.excel}")
    if not os.path.isfile(args.zip):
        sys.exit(f"ZIP file not found: {args.zip}")
    if do_live and not args.base_url:
        sys.exit("--base-url is required unless you pass --static-only")

    selector_map = None
    if args.selector_map:
        if not os.path.isfile(args.selector_map):
            sys.exit(f"--selector-map file not found: {args.selector_map}")
        with open(args.selector_map, "r", encoding="utf-8") as fh:
            selector_map = json.load(fh)

    pw_module = None
    if do_live:
        try:
            from playwright.sync_api import sync_playwright as _sp
            pw_module = _sp
        except ImportError:
            sys.exit("This script needs Playwright for the live run. Install it with:\n"
                      "    pip install playwright --break-system-packages\n"
                      "    playwright install chromium\n"
                      "(or pass --static-only to skip the live run and just check source code)")

    output_path = args.output or (os.path.splitext(args.excel)[0] + "_COMBINED_RESULTS.xlsx")
    mode = "static + live" if (do_static and do_live) else ("static only" if do_static else "live only")

    print(f"[1/7] Extracting {args.zip} ...")
    tmpdir = extract_zip(args.zip)

    proc = logf = None
    index = None
    n_files = None
    try:
        if do_static:
            print("[2/7] Indexing source code for the static check ...")
            index = CodebaseIndex(tmpdir)
            n_files = len(index.files)
            print(f"      indexed {n_files} text files")

        if do_live:
            if args.start_cmd:
                print("[3/7] Starting the app ...")
                cwd = os.path.join(tmpdir, args.cwd) if args.cwd else tmpdir
                log_path = os.path.join(tmpdir, "app_startup.log")
                proc, logf = start_app(args.start_cmd, cwd, log_path)
                if not wait_for_server(args.base_url, args.startup_timeout):
                    print(f"      see {log_path} for what the app printed on startup")
                    sys.exit("App never became reachable at --base-url.")
                print(f"      app is up at {args.base_url}")
            else:
                print("[3/7] Checking the app is already reachable ...")
                if not wait_for_server(args.base_url, 5):
                    sys.exit("Nothing responded at --base-url. Start the app first, or pass --start-cmd.")
        else:
            print("[3/7] Skipping live run (--static-only)")

        print(f"[4/7] Loading test cases from sheet '{args.sheet}' ...")
        wb = openpyxl.load_workbook(args.excel)
        if args.sheet not in wb.sheetnames:
            sys.exit(f"Sheet '{args.sheet}' not found. Sheets: {wb.sheetnames}")
        ws = wb[args.sheet]
        col_of = load_header_map(ws)

        full_rows = []
        for r in range(2, ws.max_row + 1):
            row_vals = {h: ws.cell(row=r, column=c).value for h, c in col_of.items()}
            if row_vals.get("Test Case ID"):
                full_rows.append(row_vals)

        creds = {}
        if do_live:
            creds = build_credentials(full_rows)
            print(f"      learned login credentials for role(s): {', '.join(sorted(creds)) or '(none found)'}")

        all_rows, row_numbers = [], []
        for r in range(2, ws.max_row + 1):
            row_vals = {h: ws.cell(row=r, column=c).value for h, c in col_of.items()}
            if not row_vals.get("Test Case ID"):
                continue
            if args.filter_module and (row_vals.get("Module") or "").lower() != args.filter_module.lower():
                continue
            all_rows.append(row_vals)
            row_numbers.append(r)
            if args.max_tests and len(all_rows) >= args.max_tests:
                break
        print(f"      {len(all_rows)} test case(s) selected to run")

        print(f"[5/7] Running checks ({mode}) ...")
        counts = Counter()
        changed = 0
        evidence_rows = []

        def process(browser_or_none):
            nonlocal changed
            for i, (r, row_vals) in enumerate(zip(row_numbers, all_rows), start=1):
                static_status = static_evidence = None
                if do_static:
                    static_status, static_evidence = evaluate_static(index, row_vals)

                live_status = live_summary = live_notes = screenshot = None
                if do_live:
                    live_result = run_live(browser_or_none, args.base_url, row_vals, creds,
                                            args.per_test_timeout, args.screenshots_dir,
                                            settle_ms=args.settle_ms, selector_map=selector_map,
                                            debug=args.debug)
                    live_status = live_result["status"]
                    live_summary = live_result["steps_summary"]
                    live_notes = live_result["notes"]
                    screenshot = live_result["screenshot"]

                final_status, confidence, combo_note = combine_verdicts(static_status, live_status)
                counts[final_status] += 1

                previous_status = row_vals.get("Status")
                if str(previous_status).strip().upper() != final_status:
                    changed += 1

                status_col = col_of["Status"]
                cell = ws.cell(row=r, column=status_col, value=final_status)
                cell.fill = STATUS_FILLS.get(final_status)

                evidence_rows.append([
                    row_vals.get("Test Case ID"), previous_status, final_status, confidence,
                    static_status or "n/a", static_evidence or "n/a",
                    live_status or "n/a", live_summary or "n/a",
                    "; ".join(x for x in [live_notes, combo_note] if x) or combo_note,
                    screenshot or "",
                ])

                if i % 10 == 0 or i == len(all_rows):
                    print(f"      {i}/{len(all_rows)} run... "
                          f"(PASS {counts['PASS']} / REVIEW {counts['REVIEW']} / FAIL {counts['FAIL']})")

        if do_live:
            with pw_module() as pw:
                browser = pw.chromium.launch(headless=not args.headed)
                try:
                    process(browser)
                finally:
                    browser.close()
        else:
            process(None)

        print("[6/7] Writing evidence log + summary ...")
        write_evidence_log(wb, evidence_rows)
        write_summary_block(wb, counts, changed, len(all_rows), args.zip, args.base_url, mode, n_files)

        print("[7/7] Saving workbook ...")
        wb.save(output_path)
    finally:
        if proc and not args.keep_server:
            print("Stopping the app ...")
            stop_app(proc, logf)
        shutil.rmtree(tmpdir, ignore_errors=True)

    print("\n=== COMBINED VALIDATION SUMMARY ===")
    for v in ["PASS", "REVIEW", "FAIL"]:
        print(f"  Status = {v:<8}: {counts.get(v, 0)}")
    print(f"  Status changed from original : {changed}")
    print(f"\nSaved: {output_path}")
    print(f"Per-row reasoning + confidence: '{EVIDENCE_SHEET_NAME}' sheet inside the workbook")


if __name__ == "__main__":
    main()